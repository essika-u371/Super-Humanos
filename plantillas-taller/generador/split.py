import sys, re, docx, copy
from docx.oxml.ns import qn

def text(el): return ''.join(x.text or '' for x in el.iter() if x.tag==qn('w:t'))

def replace_text(doc, old, new):
    """Reemplaza texto en párrafos del cuerpo, encabezados y pies, run por run o uniendo runs."""
    parts=[doc.element.body]
    for s in doc.sections:
        for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            try: parts.append(hf._element)
            except Exception: pass
    n=0
    for part in parts:
        for p in part.iter(qn('w:p')):
            ts=[t for t in p.iter(qn('w:t'))]
            full=''.join(t.text or '' for t in ts)
            if old not in full:
                if old.upper() in full: old_, new_ = old.upper(), new.upper()
                else: continue
            else: old_, new_ = old, new
            hit=False
            for t in ts:
                if t.text and old_ in t.text: t.text=t.text.replace(old_,new_); hit=True; n+=1
            if not hit:
                full=full.replace(old_,new_); ts[0].text=full
                for t in ts[1:]: t.text=''
                n+=1
    return n

def body_index(doc, starts):
    body=list(doc.element.body)
    for i,el in enumerate(body):
        if el.tag==qn('w:p') and text(el).strip().upper().startswith(starts.upper()): return i
    raise SystemExit(f'no encontré {starts!r}')

def remove_range(doc, a_txt, b_txt):
    body=doc.element.body; els=list(body)
    a=body_index(doc,a_txt); b=body_index(doc,b_txt) if b_txt else len(els)-1
    for el in els[a:b]: body.remove(el)

def table_label(t):
    return {text(r[1]) if False else None for r in []}

OTHER={'comercial':'Administrativo','administrativo':'Comercial'}
OWN={'comercial':'Comercial','administrativo':'Administrativo'}

def keep_tag(val, perfil):
    v=val.strip()
    return v.startswith(OWN[perfil]) or v.startswith('Los dos perfiles') or not v.startswith(OTHER[perfil])

def temario(src, dst, perfil):
    d=docx.Document(src)
    if perfil=='comercial':
        remove_range(d,'6 · TEMARIO DEL PERFIL ADMINISTRATIVO','7 · LOS RETOS DE LOS JUEVES')
        replace_text(d,'7 · LOS RETOS DE LOS JUEVES','6 · LOS RETOS DE LOS JUEVES')
    else:
        remove_range(d,'5 · TEMARIO DEL PERFIL COMERCIAL','6 · TEMARIO DEL PERFIL ADMINISTRATIVO')
        replace_text(d,'6 · TEMARIO DEL PERFIL ADMINISTRATIVO','5 · TEMARIO DEL PERFIL ADMINISTRATIVO')
        replace_text(d,'7 · LOS RETOS DE LOS JUEVES','6 · LOS RETOS DE LOS JUEVES')
    # tabla de casos de la sección 4: quitar filas del otro perfil
    for t in d.tables:
        if len(t.columns)==2 and any(OTHER[perfil]+' · taller' in c.text for r in t.rows for c in r.cells):
            for r in list(t.rows)[1:]:
                if not keep_tag(r.cells[1].text, perfil): t._tbl.remove(r._tr)
    replace_text(d,'Temario detallado de los 12 talleres',f'Temario detallado · perfil {perfil}')
    replace_text(d,'Perfil comercial y perfil administrativo · guía para el docente',f'Los seis talleres del perfil {perfil} · guía para el docente')
    d.save(dst)

def guia(src, dst, perfil):
    d=docx.Document(src)
    if perfil=='comercial':
        remove_range(d,'3 · TALLER 1 · ADMINISTRATIVO','4 · LO QUE DEBE CUBRIR')
        remove_range(d,'5 · LO QUE DEBE CUBRIR','6 · ESTÁNDARES')
        replace_text(d,'4 · LO QUE DEBE CUBRIR','3 · LO QUE DEBE CUBRIR'); replace_text(d,'6 · ESTÁNDARES','4 · ESTÁNDARES')
        replace_text(d,'9 y 16 de octubre','23 y 30 de octubre')
    else:
        remove_range(d,'2 · TALLER 1 · COMERCIAL','3 · TALLER 1 · ADMINISTRATIVO')
        remove_range(d,'4 · LO QUE DEBE CUBRIR','5 · LO QUE DEBE CUBRIR')
        replace_text(d,'3 · TALLER 1 · ADMINISTRATIVO','2 · TALLER 1 · ADMINISTRATIVO')
        replace_text(d,'5 · LO QUE DEBE CUBRIR','3 · LO QUE DEBE CUBRIR'); replace_text(d,'6 · ESTÁNDARES','4 · ESTÁNDARES')
        replace_text(d,'14 y 21 de octubre','23 y 28 de octubre')
    replace_text(d,'Los doce talleres tienen','Los seis talleres tienen')
    replace_text(d,'Guía de contenido por taller',f'Guía de contenido · perfil {perfil}')
    d.save(dst)

def casos(src, dst, perfil):
    d=docx.Document(src); body=d.element.body; els=list(body)
    removed=0
    for i,el in enumerate(els):
        if el.tag!=qn('w:tbl'): continue
        rows=el.findall(qn('w:tr'))
        lab={}
        for tr in rows:
            tcs=tr.findall(qn('w:tc'))
            if len(tcs)>=2: lab[text(tcs[0]).strip()]=text(tcs[1]).strip()
        v=lab.get('Dónde se usa')
        if v is None or keep_tag(v, perfil): continue
        # retroceder hasta el párrafo «CASO n»
        j=i
        while j>0 and not text(els[j]).strip().startswith('CASO'): j-=1
        k=i+1
        while k<len(els) and els[k].tag==qn('w:p') and not text(els[k]).strip(): k+=1
        for el2 in els[j:k]:
            if el2.getparent() is not None: body.remove(el2)
        removed+=1
    left=sum(1 for el in body if el.tag==qn('w:p') and text(el).strip().startswith('CASO'))
    replace_text(d,'Material de clase para los docentes · 24 casos reales levantados en las mesas de trabajo',f'Material de clase para el docente del perfil {perfil} · {left} casos reales levantados en las mesas de trabajo')
    replace_text(d,'Guía de casos de G&J',f'Guía de casos de G&J · perfil {perfil}')
    d.save(dst); print('casos', perfil, 'quitados', removed, 'quedan', left)

def ajustes(src, dst, perfil):
    d=docx.Document(src)
    if perfil=='comercial':
        remove_range(d,'5 · Ajustes por taller · administrativo','6 · Retos de los jueves')
    else:
        remove_range(d,'4 · Ajustes por taller · comercial','5 · Ajustes por taller · administrativo')
        replace_text(d,'5 · Ajustes por taller · administrativo','4 · Ajustes por taller · administrativo')
    remove_range(d,'7 · Decisiones pendientes',None)
    replace_text(d,'6 · Retos de los jueves','5 · Retos de los jueves')
    replace_text(d,'3 · Ajustes que aplican a los doce talleres','3 · Ajustes que aplican a los seis talleres')
    replace_text(d,'documento de trabajo para la dirección del programa y los dos docentes',f'documento para el docente del perfil {perfil}')
    replace_text(d,'AJUSTES A LOS TALLERES · DIAGNÓSTICO Y POLÍTICA DE IA',f'AJUSTES A LOS TALLERES · PERFIL {perfil.upper()}')
    d.save(dst)

if __name__=='__main__':
    kind, src, dst, perfil = sys.argv[1:5]
    {'temario':temario,'guia':guia,'casos':casos,'ajustes':ajustes}[kind](src,dst,perfil)
